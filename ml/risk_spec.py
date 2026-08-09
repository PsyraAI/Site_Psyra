"""Leitura auditável do gabarito XLSX e fórmula de risco 0–100."""

from __future__ import annotations

import json
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook


@dataclass(frozen=True)
class ItemRisco:
    """Regra de pontuação de um item do questionário."""

    codigo: str
    bloco: str
    dimensao: str
    peso: float
    texto: str
    reverso: bool
    validacao_crp: str | None


def carregar_gabarito_xlsx(caminho: Path) -> list[ItemRisco]:
    """Lê a aba Gabarito sem depender de fórmulas calculadas pelo Excel."""
    workbook = load_workbook(caminho, data_only=False, read_only=True)
    if "Gabarito" not in workbook.sheetnames:
        raise ValueError("aba 'Gabarito' ausente no XLSX")
    aba = workbook["Gabarito"]

    itens: list[ItemRisco] = []
    for linha in range(4, aba.max_row + 1):
        codigo = aba.cell(linha, 1).value
        if not codigo:
            continue
        direcao = str(aba.cell(linha, 6).value or "").strip().casefold()
        if direcao not in {"protetor", "risco"}:
            raise ValueError(f"direção inválida para {codigo}: {direcao!r}")
        validacao = aba.cell(linha, 14).value
        itens.append(
            ItemRisco(
                codigo=str(codigo).strip(),
                bloco=str(aba.cell(linha, 2).value).strip(),
                dimensao=str(aba.cell(linha, 3).value).strip(),
                peso=float(aba.cell(linha, 4).value),
                texto=str(aba.cell(linha, 5).value).strip(),
                reverso=direcao == "protetor",
                validacao_crp=str(validacao).strip() if validacao else None,
            )
        )
    if len(itens) != 46:
        raise ValueError(f"esperados 46 itens no XLSX; encontrados {len(itens)}")
    if len({item.codigo for item in itens}) != len(itens):
        raise ValueError("códigos duplicados no XLSX")
    return itens


def carregar_instrumento_json(caminho: Path) -> dict[str, Any]:
    """Carrega o instrumento operacional usado pelo backend."""
    with caminho.open(encoding="utf-8") as arquivo:
        return json.load(arquivo)


def validar_xlsx_contra_json(
    itens: list[ItemRisco], instrumento: dict[str, Any]
) -> dict[str, Any]:
    """Falha se código, bloco, texto, peso ou polaridade divergirem."""
    esperados: dict[str, dict[str, Any]] = {}
    for bloco in instrumento["blocos"]:
        codigo_bloco = str(bloco["codigo"])
        reverso_bloco = bool(bloco.get("reverso", False))
        for indice, item in enumerate(bloco["itens"], start=1):
            codigo = f"{codigo_bloco}{indice}"
            texto = str(item.get("texto", "")) if isinstance(item, dict) else str(item)
            reverso = (
                bool(item.get("reverso", reverso_bloco))
                if isinstance(item, dict)
                else reverso_bloco
            )
            esperados[codigo] = {
                "bloco": codigo_bloco,
                "dimensao": str(bloco["dimensao"]),
                "peso": float(bloco.get("peso", 1.0)),
                "texto": texto.strip(),
                "reverso": reverso,
            }

    divergencias: list[str] = []
    recebidos = {item.codigo: item for item in itens}
    if set(recebidos) != set(esperados):
        divergencias.append("conjunto de códigos difere entre XLSX e JSON")
    for codigo in sorted(set(recebidos).intersection(esperados)):
        recebido = asdict(recebidos[codigo])
        for campo, esperado in esperados[codigo].items():
            if recebido[campo] != esperado:
                divergencias.append(
                    f"{codigo}.{campo}: XLSX={recebido[campo]!r}, JSON={esperado!r}"
                )
    if divergencias:
        raise ValueError("gabarito incompatível:\n" + "\n".join(divergencias))

    return {
        "n_itens": len(itens),
        "n_protetores": sum(item.reverso for item in itens),
        "n_risco": sum(not item.reverso for item in itens),
        "n_validacoes_crp": sum(item.validacao_crp is not None for item in itens),
        "blocos": sorted({item.bloco for item in itens}),
    }


def orientar_itens_risco(respostas: pd.DataFrame, itens: list[ItemRisco]) -> pd.DataFrame:
    """Converte Likert 1–5 para direção única: maior valor = maior risco."""
    ausentes = sorted({item.codigo for item in itens} - set(respostas.columns))
    if ausentes:
        raise ValueError(f"itens ausentes: {', '.join(ausentes)}")

    orientados: dict[str, pd.Series] = {}
    for item in itens:
        valor = pd.to_numeric(respostas[item.codigo], errors="coerce")
        invalidos = valor.notna() & ~valor.between(1, 5)
        if invalidos.any():
            raise ValueError(f"resposta fora da escala 1–5 em {item.codigo}")
        orientados[item.codigo] = 6.0 - valor if item.reverso else valor
    return pd.DataFrame(orientados, index=respostas.index, dtype=float)


def calcular_indice_risco(respostas: pd.DataFrame, itens: list[ItemRisco]) -> pd.Series:
    """Aplica exatamente a média ponderada por bloco declarada no XLSX."""
    orientados = orientar_itens_risco(respostas, itens)
    por_bloco: dict[str, list[str]] = {}
    pesos: dict[str, float] = {}
    for item in itens:
        por_bloco.setdefault(item.bloco, []).append(item.codigo)
        pesos[item.bloco] = item.peso

    numerador = pd.Series(0.0, index=respostas.index)
    denominador = pd.Series(0.0, index=respostas.index)
    for bloco, colunas in por_bloco.items():
        media = orientados[colunas].mean(axis=1)
        indice_bloco = (media - 1.0) / 4.0 * 100.0
        valido = media.notna()
        numerador = numerador.add(indice_bloco.fillna(0.0) * pesos[bloco])
        denominador = denominador.add(valido.astype(float) * pesos[bloco])
    return numerador / denominador.replace(0.0, float("nan"))
